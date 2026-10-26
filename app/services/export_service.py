import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.consultation import Consultation, ConsultationStatus

SECTIONS = [
    ("complaints", "Жалобы"),
    ("anamnesis", "Анамнез"),
    ("symptoms_duration", "Симптомы и их продолжительность"),
    ("examination", "Объективные данные"),
    ("diagnosis", "Заключение (диагноз)"),
    ("treatment_plan", "План лечения"),
    ("recommendations", "Рекомендации"),
]
_LABELS = {
    "ru": {
        "complaints": "Жалобы", "anamnesis": "Анамнез", "symptoms_duration": "Симптомы и их продолжительность",
        "examination": "Объективные данные", "diagnosis": "Заключение (диагноз)",
        "treatment_plan": "План лечения", "recommendations": "Рекомендации",
        "title": "Медицинское заключение", "patient": "Пациент", "doctor": "Врач",
        "date": "Дата приёма", "confirmed": "Подтверждено врачом", "draft": "ЧЕРНОВИК: не подтверждён врачом",
        "missing": "Требует уточнения", "memo": "ПАМЯТКА ПАЦИЕНТУ",
        "disclaimer": "Документ подготовлен с помощью AqylMed AI и должен быть проверен врачом.",
    },
    "kk": {
        "complaints": "Шағымдар", "anamnesis": "Анамнез", "symptoms_duration": "Белгілердің ұзақтығы",
        "examination": "Объективті деректер", "diagnosis": "Қорытынды (диагноз)",
        "treatment_plan": "Емдеу жоспары", "recommendations": "Ұсыныстар",
        "title": "Медициналық қорытынды", "patient": "Пациент", "doctor": "Дәрігер",
        "date": "Қабылдау күні", "confirmed": "Дәрігер растады", "draft": "ЖОБА: дәрігер растамаған",
        "missing": "Нақтылауды қажет етеді", "memo": "Пациентке арналған жадынама",
        "disclaimer": "AqylMed AI көмегімен әзірленген. Дәрігер тексеруі керек.",
    },
    "en": {
        "complaints": "Chief complaints", "anamnesis": "History", "symptoms_duration": "Symptoms and duration",
        "examination": "Examination", "diagnosis": "Assessment (diagnosis)",
        "treatment_plan": "Treatment plan", "recommendations": "Recommendations",
        "title": "Medical note", "patient": "Patient", "doctor": "Doctor",
        "date": "Visit date", "confirmed": "Confirmed by doctor", "draft": "DRAFT: not confirmed by doctor",
        "missing": "Needs clarification", "memo": "Patient instructions",
        "disclaimer": "Prepared with AqylMed AI. Must be reviewed by a doctor.",
    },
}


def _label(key: str, locale: str, custom: dict[str, str] | None = None) -> str:
    return (custom or {}).get(key) or _LABELS.get(locale, {}).get(key) or dict(SECTIONS).get(key, key)


def _status_line(c: Consultation, locale: str = "ru", labels: dict[str, str] | None = None) -> str:
    confirmed_label = _label("confirmed", locale, labels)
    draft_label = _label("draft", locale, labels)
    if c.status == ConsultationStatus.CONFIRMED and c.confirmed_at:
        who = c.confirmed_by.full_name if c.confirmed_by else c.doctor_name
        return f"{confirmed_label}: {who}, {c.confirmed_at:%d.%m.%Y %H:%M}"
    return draft_label


def _missing(doc) -> list[str]:
    return [line for line in (doc.missing_information or "").splitlines() if line.strip()]


def export_txt(c: Consultation, locale: str = "ru", field_labels: dict[str, str] | None = None) -> str:
    titles = {
        "ru": {"initial": "Первичный приём", "examination": "Осмотр", "discharge": "Выписка"},
        "kk": {"initial": "Алғашқы қабылдау", "examination": "Тексеру", "discharge": "Үзінді көшірме"},
        "en": {"initial": "Initial consultation", "examination": "Examination", "discharge": "Discharge summary"},
    }
    field_labels = {**(field_labels or {}), "title": _label("title", locale) + " — " + titles.get(locale, titles["ru"]).get(c.document_type, "")}
    label = lambda key: _label(key, locale, field_labels)
    doc, memo = c.medical_document, c.patient_memo
    lines = [
        label("title").upper(),
        f"{label('patient')}: {c.patient_name}",
        f"{label('doctor')}: {c.doctor_name}",
        f"{label('date')}: {c.consultation_date}",
        _status_line(c, locale, field_labels),
        "",
    ]
    if doc:
        for key, title in SECTIONS:
            content = getattr(doc, key) or ("көрсетілмеген" if locale == "kk" else "not specified" if locale == "en" else "не указано")
            if key == "diagnosis" and doc.diagnosis_code:
                content = f"{doc.diagnosis_code} — {content}"
            lines += [f"{label(key)}:", content, ""]
        if _missing(doc):
            lines += [f"{label('missing')}:"] + [f"- {m}" for m in _missing(doc)] + [""]
    if memo:
        lines += [label("memo").upper(), memo.memo_text, ""]
    lines.append(label("disclaimer"))
    return "\n".join(lines)


def export_docx(c: Consultation, locale: str = "ru", field_labels: dict[str, str] | None = None) -> bytes:
    titles = {
        "ru": {"initial": "Первичный приём", "examination": "Осмотр", "discharge": "Выписка"},
        "kk": {"initial": "Алғашқы қабылдау", "examination": "Тексеру", "discharge": "Үзінді көшірме"},
        "en": {"initial": "Initial consultation", "examination": "Examination", "discharge": "Discharge summary"},
    }
    field_labels = {**(field_labels or {}), "title": _label("title", locale) + " — " + titles.get(locale, titles["ru"]).get(c.document_type, "")}
    label = lambda key: _label(key, locale, field_labels)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = document.add_heading(label("title"), level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for field, value in ((label("patient"), c.patient_name), (label("doctor"), c.doctor_name),
                         (label("date"), c.consultation_date)):
        p = document.add_paragraph()
        p.add_run(f"{field}: ").bold = True
        p.add_run(value)
    status = document.add_paragraph()
    run = status.add_run(_status_line(c, locale, field_labels))
    run.bold = True
    run.italic = c.status != ConsultationStatus.CONFIRMED

    doc, memo = c.medical_document, c.patient_memo
    if doc:
        for key, title_text in SECTIONS:
            document.add_heading(label(key), level=2)
            content = getattr(doc, key) or ("көрсетілмеген" if locale == "kk" else "not specified" if locale == "en" else "не указано")
            if key == "diagnosis" and doc.diagnosis_code:
                content = f"{doc.diagnosis_code} — {content}"
            document.add_paragraph(content)
        if _missing(doc):
            document.add_heading(label("missing"), level=2)
            for item in _missing(doc):
                document.add_paragraph(item, style="List Bullet")
    if memo:
        document.add_page_break()
        document.add_heading(label("memo"), level=1)
        for line in memo.memo_text.splitlines():
            document.add_paragraph(line)

    footer = document.add_paragraph()
    note = footer.add_run(label("disclaimer"))
    note.italic = True
    note.font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
